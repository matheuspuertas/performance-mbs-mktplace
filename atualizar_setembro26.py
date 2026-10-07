"""
Atualiza a planilha Performance MBS com os dados de Setembro/2026.
Executa sem interacao - todos os valores foram extraidos e verificados.

Grava: Mercado Livre e Shopee (completos).
Site (parcial): numero de pedidos e valor total, do CSV de pedidos da Tray
         (42 pedidos, canal "LOJA VIRTUAL", nenhum cancelado). Ticket medio, valor
         investido, TACOS e conversao sao formulas que ja estao na linha da planilha.
         Faltam: usuarios, Google Ads e Meta Ads (ainda nao enviados pelo usuario).

ATENCAO (Shopee publicidade): a pasta tem dois CSVs de anuncios com periodo 06/09 a 06/10
         ("Ultimo mes" no painel = ultimos 30 dias) e um com o periodo certo, 01/09 a 30/09.
         Este script le o CSV certo pelo nome (SHOPEE_ADS_CSV), sem depender de
         encontrar_arquivos(). Conferido com a captura 210233.png (01/09 - 30/09):
         vendas R$209.100,39, investimento R$12.050,17, ROAS 17,35 -- batem.
         A captura 205146.png e do periodo 06/09 a 06/10 e foi ignorada.

A linha de setembro da aba Shopee tinha valores digitados a mao em Receita ADS Direto,
Despesas e ROAS (200000 / 12000 / 17); o usuario autorizou apagar -- sao substituidos
pelos valores do CSV.

ATENCAO: o relatorio "Relatorio_evolucao_negocio_2026_09_01-2026_09_30.xlsx" voltou a vir
         com uma linha por dia (diferente de jul/ago, que vinham com uma linha por mes),
         mas traz tambem o dia 01/10/26 (vazamento do mes seguinte). Por isso este script
         usa processar_ml_evolucao_dias(), que soma apenas as linhas de 09/2026.
         Conferido com a captura "Resumo de desempenho" (192425.png): visitas 73.923,
         quantidade de vendas 5.238, unidades 5.525, vendas brutas R$768.014 -- batem.

Imagens ignoradas:
  - Captura de tela 2026-10-06 190853.png e 191235.png -> painel de Publicidade do ML
    com periodo 1-6 de outubro, fora do mes de referencia.
"""
import sys
sys.path.insert(0, r"P:\Meu Drive\Empresas\MBS Pro Grooming\Desenvolvimento de relatórios\projeto 1")

import openpyxl
import pandas as pd
import shutil
from pathlib import Path

from update_performance import (
    PERFORMANCE_FILE,
    MATERIAIS_DIR,
    ML_COLS, ML_FORMULA_COLS, ML_LINHA_DADOS,
    SHOPEE_COLS, SHOPEE_FORMULA_COLS, SHOPEE_LINHA_DADOS,
    SITE_COLS, SITE_FORMULA_COLS, SITE_LINHA_DADOS,
    encontrar_arquivos,
    encontrar_linha_mes,
    processar_shopee_overview,
    processar_shopee_ads_csv,
    mes_serial_excel,
    atualizar_aba,
    limpar_numero,
)

from graficos import restaurar_graficos

ANO = 2026
MES = 9

TRAY_CSV = Path(MATERIAIS_DIR) / "pedidos_1277450_3168802a-4202-4eea-a62d-54be2ed9b0a1.csv"
SHOPEE_ADS_CSV = Path(MATERIAIS_DIR) / "Dados+Gerais+de+Anúncios+Shopee-01_09_2026-30_09_2026.csv"


def processar_ml_evolucao_dias(caminho, ano, mes):
    """Soma apenas as linhas diarias (dd/mm/aa) do mes/ano informado no relatorio de
    evolucao do ML, ignorando dias de outros meses (ex: vazamento do mes seguinte)."""
    print(f"\n  >> Lendo relatorio ML: {caminho.name}  (somente dias de {mes:02d}/{ano})")
    wb = openpyxl.load_workbook(str(caminho), data_only=True)

    aba = None
    for nome in wb.sheetnames:
        if "neg" in nome.lower():
            aba = wb[nome]
            break
    if aba is None:
        aba = wb.active

    header_row = None
    for r in range(1, 20):
        val = aba.cell(r, 1).value
        if val and str(val).strip().lower() == "data":
            header_row = r
            break
    if header_row is None:
        raise ValueError("Cabecalho 'Data' nao encontrado na aba Negocio.")

    col_map = {}
    for c in range(1, aba.max_column + 1):
        val = aba.cell(header_row, c).value
        if val:
            col_map[str(val).strip()] = c

    campos = {
        "Visitas":                "visitas",
        "Compradores únicos":     "compradores_unicos",
        "Novos compradores":      "novos_compradores",
        "Compradores existentes": "compradores_existentes",
        "Quantidade de vendas":   "qtd_vendas",
        "Unidades vendidas":      "unidades_vendidas",
        "Vendas brutas":          "vendas_brutas",
    }

    sufixo = f"/{mes:02d}/{str(ano)[2:]}"
    totais = {v: 0.0 for v in campos.values()}
    linhas = 0
    ignoradas = []

    for r in range(header_row + 1, aba.max_row + 1):
        data_val = aba.cell(r, 1).value
        if not data_val or not str(data_val).strip():
            continue
        data_txt = str(data_val).strip()
        if not data_txt.endswith(sufixo):
            ignoradas.append(data_txt)
            continue
        for campo_orig, campo_dest in campos.items():
            col = col_map.get(campo_orig)
            val = aba.cell(r, col).value if col else None
            if isinstance(val, str):
                val = limpar_numero(val)
            if val is not None:
                totais[campo_dest] += float(val)
        linhas += 1

    if linhas == 0:
        raise ValueError(f"Nenhuma linha de {mes:02d}/{ano} encontrada no relatorio.")

    cu, vb, qv, uv, vi, ce = (
        totais["compradores_unicos"], totais["vendas_brutas"], totais["qtd_vendas"],
        totais["unidades_vendidas"], totais["visitas"], totais["compradores_existentes"],
    )
    totais["media_vendas_comprador"] = vb / cu if cu else 0
    totais["taxa_recompra"]          = ce / cu if cu else 0
    totais["conversao"]              = qv / vi if vi else 0
    totais["valor_medio_venda"]      = vb / qv if qv else 0
    totais["preco_medio_unidade"]    = vb / uv if uv else 0

    print(f"  [OK] {linhas} dias agregados: {vi:,.0f} visitas, R${vb:,.2f} vendas brutas")
    if ignoradas:
        print(f"  [!] Linhas ignoradas (fora do mes): {', '.join(ignoradas)}")
    return totais

# -------------------------------------------------------------------------------
# DADOS DAS IMAGENS (verificados manualmente)
# -------------------------------------------------------------------------------

# Imagens: Captura de tela 2026-10-06 190727.png + 190743.png (ML Publicidade - Metricas ao vivo, 1-30 set)
ML_ADS = {
    "pub_impressoes":          9_235_365,
    "pub_cliques":             25_849,
    "pub_vendas_product_ads":  2_769,     # "Vendas atribuidas"
    "pub_vendas_sem_products": 1_543,     # "Outras vendas"
    "pub_investimento":        23_913.00,
    "pub_receita":             409_036.00,
}

# Imagem: Captura de tela 2026-10-06 192332.png (ML Afiliados - todas as vendas de afiliados, 1-30 set)
ML_AFILIADOS = {
    "af_receita":    72_590.22,
    "af_unidades":   476,
    "af_qtd_vendas": 448,
    "af_custo":      3_879.54,
}

# Imagem: Captura de tela 2026-10-06 205131.png (Shopee Afiliados - Principais Indicadores, 09/2026)
SHOPEE_AFILIADOS = {
    "af_vendas":             72_800.00,
    "af_itens_brutos":       685,
    "af_pedidos":            564,
    "af_cliques":            24_800,
    "af_comissao":           4_900.00,
    "af_roi":                15,
    "af_compradores_totais": 546,
    "af_novos_compradores":  353,
}

# -------------------------------------------------------------------------------
# EXECUCAO
# -------------------------------------------------------------------------------

print("=" * 62)
print("   ATUALIZACAO SETEMBRO/2026 -- PERFORMANCE MBS MKTPLACE")
print("=" * 62)

arqs = encontrar_arquivos(MATERIAIS_DIR)

print("\n[1/3] Lendo relatorio ML evolucao...")
dados_ml = processar_ml_evolucao_dias(arqs["ml_evolucao"], ANO, MES)
dados_ml.update(ML_ADS)
dados_ml.update(ML_AFILIADOS)

print("\n[2/3] Lendo Shopee overview + CSV...")
dados_shopee = processar_shopee_overview(arqs["shopee_overview"])
dados_shopee.update(processar_shopee_ads_csv(SHOPEE_ADS_CSV))
dados_shopee.update(SHOPEE_AFILIADOS)

print("\n[3/3] Lendo pedidos do site (Tray)...")
df_tray = pd.read_csv(str(TRAY_CSV), encoding="latin1", sep=None, engine="python")
total_tray = (
    df_tray["Total"].astype(str).str.replace(".", "", regex=False).str.replace(",", ".", regex=False).astype(float)
)
# So os campos de entrada: as colunas calculadas da linha sao formulas e nao sao tocadas.
dados_site = {
    "num_pedidos": len(df_tray),
    "valor_total": round(float(total_tray.sum()), 2),
}
print(f"  [OK] {dados_site['num_pedidos']} pedidos | Valor total: R${dados_site['valor_total']:,.2f}")

print("\n" + "=" * 62)
print("  RESUMO DO QUE SERA GRAVADO -- set/26")
print("=" * 62)

print("\nMercado Livre:")
print(f"  Visitas:             {dados_ml['visitas']:,.0f}")
print(f"  Compradores unicos:  {dados_ml['compradores_unicos']:,.0f}")
print(f"  Novos compradores:   {dados_ml['novos_compradores']:,.0f}")
print(f"  Compr. existentes:   {dados_ml['compradores_existentes']:,.0f}")
print(f"  Qtd vendas:          {dados_ml['qtd_vendas']:,.0f}")
print(f"  Unidades vendidas:   {dados_ml['unidades_vendidas']:,.0f}")
print(f"  Vendas brutas:       R${dados_ml['vendas_brutas']:,.2f}")
print(f"  Conversao:           {dados_ml['conversao']*100:.2f}%")
print(f"  Impressoes ads:      {dados_ml['pub_impressoes']:,.0f}")
print(f"  Cliques ads:         {dados_ml['pub_cliques']:,.0f}")
print(f"  Investimento ads:    R${dados_ml['pub_investimento']:,.2f}")
print(f"  Receita ads:         R${dados_ml['pub_receita']:,.2f}")
print(f"  Rec. afiliados:      R${dados_ml['af_receita']:,.2f}")
print(f"  Unid. afiliados:     {dados_ml['af_unidades']:,.0f}")
print(f"  Qtd vendas afil.:    {dados_ml['af_qtd_vendas']:,.0f}")
print(f"  Custo afiliados:     R${dados_ml['af_custo']:,.2f}")

print("\nShopee:")
print(f"  Visitantes:          {dados_shopee['visitantes']:,.0f}")
print(f"  Compradores feitos:  {dados_shopee['num_pedidos']:,.0f}")
print(f"  Vendas feitos:       R${dados_shopee['pedidos_feitos_valor']:,.2f}")
print(f"  Compradores pagos:   {dados_shopee['pedidos_pagos_qtd']:,.0f}")
print(f"  Vendas pagos:        R${dados_shopee['pedidos_pagos_valor']:,.2f}")
print(f"  Ticket medio:        R${dados_shopee['ticket_medio']:,.2f}")
print(f"  Impressoes ads:      {dados_shopee['pub_impressoes']:,.0f}")
print(f"  Cliques ads:         {dados_shopee['pub_cliques']:,.0f}")
print(f"  GMV ads:             R${dados_shopee['pub_gmv']:,.2f}")
print(f"  Receita direta ads:  R${dados_shopee['pub_receita_ads']:,.2f}")
print(f"  Despesas ads:        R${dados_shopee['pub_despesas']:,.2f}")
print(f"  Vendas afiliados:    R${dados_shopee['af_vendas']:,.2f}")
print(f"  Comissao afiliados:  R${dados_shopee['af_comissao']:,.2f}")

print("\nSite (parcial -- faltam usuarios, Google Ads e Meta Ads):")
print(f"  N. Pedidos:          {dados_site['num_pedidos']:,}")
print(f"  Valor Total:         R${dados_site['valor_total']:,.2f}")

# Backup
bkp = Path(PERFORMANCE_FILE).with_suffix(f".bkp_{ANO}{MES:02d}.xlsx")
if not bkp.exists():
    shutil.copy2(PERFORMANCE_FILE, bkp)
    print(f"\n[OK] Backup salvo: {bkp.name}")
else:
    print(f"\n[OK] Backup ja existe: {bkp.name}")

# Carregar e gravar
wb = openpyxl.load_workbook(PERFORMANCE_FILE)

ws = wb["Mercado Livre"]
# A planilha ja traz a linha do mes pre-criada com data na coluna A; preserva o valor
# original (atualizar_aba gravaria o serial numerico por cima).
ln_existente = encontrar_linha_mes(ws, mes_serial_excel(ANO, MES), ML_LINHA_DADOS)
mes_original = ws.cell(ln_existente, 1).value if ln_existente else None
ln = atualizar_aba(ws, ML_COLS, ML_FORMULA_COLS, dados_ml, ANO, MES, ML_LINHA_DADOS)
if mes_original is not None:
    ws.cell(ln, 1).value = mes_original
print(f"[OK] Aba Mercado Livre -> linha {ln}")

ws = wb["Shopee"]
ln_existente = encontrar_linha_mes(ws, mes_serial_excel(ANO, MES), SHOPEE_LINHA_DADOS)
mes_original = ws.cell(ln_existente, 1).value if ln_existente else None
ln = atualizar_aba(ws, SHOPEE_COLS, SHOPEE_FORMULA_COLS, dados_shopee, ANO, MES, SHOPEE_LINHA_DADOS)
if mes_original is not None:
    ws.cell(ln, 1).value = mes_original
print(f"[OK] Aba Shopee -> linha {ln}")

ws = wb["Site"]
ln_existente = encontrar_linha_mes(ws, mes_serial_excel(ANO, MES), SITE_LINHA_DADOS)
mes_original = ws.cell(ln_existente, 1).value if ln_existente else None
ln = atualizar_aba(ws, SITE_COLS, SITE_FORMULA_COLS, dados_site, ANO, MES, SITE_LINHA_DADOS)
if mes_original is not None:
    ws.cell(ln, 1).value = mes_original
print(f"[OK] Aba Site -> linha {ln}")

# Salvar (com fallback para arquivo temporario caso esteja bloqueado)
destino = Path(PERFORMANCE_FILE)
temp    = destino.parent / ("~temp_" + destino.name)

try:
    wb.save(str(destino))
    print(f"\n[OK] Planilha salva!")
    print(f"     {destino}")
    # O save do openpyxl apaga os rotulos dos graficos; restaura a partir do backup
    # (salvo pelo Excel) e acrescenta o mes novo.
    restaurar_graficos(bkp, destino, ANO, MES)
except PermissionError:
    wb.save(str(temp))
    print(f"\n[!] Arquivo bloqueado (aberto no Excel ou sincronizando no Google Drive).")
    print(f"    Arquivo salvo em: {temp.name}")
    print()
    print("    Para finalizar:")
    print("    1. Feche a planilha no Excel (se estiver aberta)")
    print("    2. Aguarde o Google Drive sincronizar")
    print(f"    3. Substitua o arquivo original pelo '{temp.name}'")
